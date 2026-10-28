"""CSI JSON/XLSX: bounded local parsing and atomic import through normalizer (T-246)."""

from __future__ import annotations

import datetime as dt
import hashlib
import io
import json
import zipfile
from dataclasses import dataclass
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from itertools import islice
from typing import Any
from xml.etree.ElementTree import ParseError
from zoneinfo import ZoneInfo

import psycopg
from defusedxml.common import DefusedXmlException
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from psycopg.rows import tuple_row
from pydantic import ValidationError

from app.contracts.events.ingest.csi_response import CsiResponse
from app.contracts.registry import validator
from app.ingest.normalizer import IngestError
from app.sim.events import dumps

MSK = ZoneInfo("Europe/Moscow")
MAX_BYTES = 10_000_000
MAX_ROWS = 10_000
# Заголовки ЕТЗ 7.2 и канонические имена схемы, без полей ФИО/контактов.
HEADERS = {
    "номер обращения": "csi_no",
    "номер csi": "csi_no",
    "номер": "csi_no",
    "текст обращения": "complaint_ru",
    "текст": "complaint_ru",
    "обращение": "complaint_ru",
    "всп": "branch_code",
    "код всп": "branch_code",
    "код талона": "ticket_code",
    "число ожидающих клиентов": "waiting_count",
    "число обслуживающихся клиентов": "serving_count",
    "число обслуженных вне очереди": "out_of_order_count",
    "время ожидания": "wait_min",
    "время ожидания (мин)": "wait_min",
    "время обслуживания": "service_min",
    "время обслуживания (мин)": "service_min",
    "число перенаправлений": "redirects",
    "оформленные продукты": "products",
    "сотрудник": "employee_public_id",
    "код сотрудника": "employee_public_id",
    "должность": "position",
    "образование": "education",
    "стаж в должности (дней)": "tenure_position_days",
    "стаж в организации (дней)": "tenure_org_days",
    "дата": "received_at",
    "дата обращения": "received_at",
    "тип обратной связи": "feedback_type",
    "тематика": "topic",
    "уточненная тематика": "subtopic",
}
INT_FIELDS = {
    "waiting_count",
    "serving_count",
    "out_of_order_count",
    "redirects",
    "tenure_position_days",
    "tenure_org_days",
}
MINUTE_FIELDS = {"wait_min", "service_min"}


class CsiError(IngestError):
    def __init__(self, column: str, message: str, row: int = 1):
        self.column = column
        self.row = row
        super().__init__(message)


@dataclass
class Row:
    number: int
    value: Any
    columns: dict[str, str]


def _header(value: Any) -> str:
    name = " ".join(str(value or "").strip().casefold().replace("ё", "е").split())
    if name in CsiResponse.model_fields:
        return name
    if name in HEADERS:
        return HEADERS[name]
    raise CsiError(str(value or "<header>")[:80], "неизвестный столбец")


def parse_json(raw: bytes) -> list[Row]:
    if len(raw) > MAX_BYTES:
        raise CsiError("<file>", "файл превышает 10 MB")
    try:
        body = json.loads(raw)
    except (ValueError, UnicodeError) as exc:
        raise CsiError("<file>", "некорректный JSON") from exc
    rows = body if isinstance(body, list) else [body]
    if not rows or len(rows) > MAX_ROWS:
        raise CsiError("<file>", "ожидается от 1 до 10000 строк")
    return [Row(i, row, {}) for i, row in enumerate(rows, start=1)]


def parse_xlsx(raw: bytes) -> list[Row]:
    if len(raw) > MAX_BYTES:
        raise CsiError("<file>", "файл превышает 10 MB")
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 50_000_000:
                raise CsiError("<file>", "распакованный XLSX превышает 50 MB")
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True, keep_links=False)
    except CsiError:
        raise
    except (ValueError, KeyError, zipfile.BadZipFile, DefusedXmlException, ParseError, InvalidFileException) as exc:
        raise CsiError("<file>", "некорректный XLSX") from exc
    try:
        sheet = workbook.worksheets[0]
        if sheet.max_column and sheet.max_column > 64:
            raise CsiError("<header>", "не больше 64 столбцов")
        cells = sheet.iter_rows(values_only=True)
        header = next(cells, None)
        if not header:
            raise CsiError("<header>", "нет заголовков")
        while header and header[-1] is None:
            header = header[:-1]
        names = [_header(value) for value in header]
        if len(set(names)) != len(names):
            raise CsiError("<header>", "повтор столбца")
        columns = dict(zip(names, map(str, header), strict=True))
        rows = []
        for number, values in enumerate(islice(cells, MAX_ROWS + 1), start=2):
            if number > MAX_ROWS + 1:
                raise CsiError("<file>", "не больше 10000 строк")
            if any(value is not None for value in values):
                if any(v is not None for v in values[len(names) :]):
                    rows.append(Row(number, {"<header>": "данные за пределами заголовка"}, columns))
                else:
                    rows.append(Row(number, dict(zip(names, values[: len(names)], strict=True)), columns))
        if not rows:
            raise CsiError("<file>", "нет строк данных")
        return rows
    finally:
        workbook.close()


def _date(value: Any) -> str:
    if isinstance(value, dt.datetime):
        stamp = value
    elif isinstance(value, dt.date):
        stamp = dt.datetime.combine(value, dt.time())
    elif isinstance(value, str):
        value = value.strip()
        try:
            stamp = dt.datetime.fromisoformat(value.replace("Z", "+00:00"))
        except ValueError:
            stamp = None
            for fmt in ("%d.%m.%Y %H:%M:%S", "%d.%m.%Y %H:%M", "%d.%m.%Y"):
                try:
                    stamp = dt.datetime.strptime(value, fmt)
                    break
                except ValueError:
                    continue
            if stamp is None:
                raise CsiError("received_at", "ожидается дата ISO 8601 либо ДД.ММ.ГГГГ") from None
    else:
        raise CsiError("received_at", "ожидается дата, а не числовой serial")
    if stamp is None:
        raise CsiError("received_at", "не удалось разобрать дату")
    return (stamp if stamp.tzinfo else stamp.replace(tzinfo=MSK)).isoformat()


def validate_row(raw: Any) -> dict[str, Any]:
    if not isinstance(raw, dict):
        raise CsiError("<row>", "ожидается объект")
    obj = {}
    for raw_key, value in raw.items():
        key = _header(raw_key)
        if key in obj:
            raise CsiError(key, "повтор столбца")
        if value is None or value == "":
            continue
        if key in INT_FIELDS | MINUTE_FIELDS:
            try:
                if isinstance(value, bool):
                    raise ValueError("boolean")
                number = Decimal(str(value).replace("\u00a0", "").replace(" ", "").replace(",", "."))
                if not number.is_finite() or number < 0:
                    raise ValueError("range")
                if key in INT_FIELDS:
                    maximum = 2147483647 if key.startswith("tenure_") else 32767
                    if number != number.to_integral_value() or number > maximum:
                        raise ValueError("integer range")
                    value = int(number)
                else:
                    if number > Decimal("999999.99"):
                        raise ValueError("minutes range")
                    value = float(number.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))
            except (ValueError, InvalidOperation) as exc:
                raise CsiError(key, "некорректное неотрицательное число или превышен диапазон") from exc
        elif key == "received_at":
            value = _date(value)
        elif key == "products" and isinstance(value, str):
            try:
                value = (
                    json.loads(value)
                    if value.lstrip().startswith("[")
                    else [v.strip() for v in value.replace(";", ",").split(",") if v.strip()]
                )
            except ValueError as exc:
                raise CsiError(key, "ожидается массив кодов продуктов") from exc
        elif isinstance(value, str):
            value = value.strip()
        obj[key] = value
    try:
        result = CsiResponse.model_validate(obj).model_dump(mode="json", exclude_none=True)
    except ValidationError as exc:
        error = exc.errors()[0]
        raise CsiError(str(error["loc"][0]), str(error["msg"])) from exc
    errors = list(validator("ingest/csi_response").iter_errors(result))
    if errors:
        error = errors[0]
        raise CsiError(str(next(iter(error.path), "<row>")), error.message)
    result.setdefault("redirects", 0)
    result.setdefault("products", [])
    return result


async def _one(conn: psycopg.AsyncConnection[Any], sql: str, args: tuple[Any, ...] = ()) -> tuple[Any, ...] | None:
    async with conn.cursor(row_factory=tuple_row) as cur:
        await cur.execute(sql, args)
        return await cur.fetchone()


async def prepare(
    conn: psycopg.AsyncConnection[Any], msg: dict[str, Any], source_code: str
) -> tuple[dict[str, Any], dict[str, Any]]:
    if source_code != "csi":
        raise CsiError("source_code", "CSI принимается источником csi")
    payload = validate_row((msg.get("meta") or {}).get("csi_meta"))
    if payload["csi_no"] != msg.get("external_id") or payload["complaint_ru"] != str(msg.get("text") or "").strip():
        raise CsiError("csi_no", "CSI metadata не совпадает с сообщением")
    await conn.execute("select pg_advisory_xact_lock(hashtext('csi_import'))")
    branch = await _one(
        conn,
        "select b.id,b.region_code,b.city_id,b.lat,b.lon,r.okato from branches b join regions r on r.code=b.region_code where b.code=%s",
        (payload["branch_code"],),
    )
    if branch is None:
        raise CsiError("branch_code", "неизвестное ВСП")
    employee = None
    if payload.get("employee_public_id"):
        employee = await _one(
            conn,
            "select id,branch_id,position,education,hired_sim,position_since_sim from employees where public_id=%s",
            (payload["employee_public_id"],),
        )
        if employee is None or employee[1] != branch[0]:
            raise CsiError("employee_public_id", "сотрудник не найден в указанном ВСП")
    detail: dict[str, Any] = {
        k: payload[k]
        for k in (
            "csi_no",
            "complaint_ru",
            "ticket_code",
            "waiting_count",
            "serving_count",
            "out_of_order_count",
            "wait_min",
            "service_min",
            "redirects",
            "products",
            "position",
            "education",
            "tenure_position_days",
            "tenure_org_days",
        )
        if k in payload
    }
    detail.update(branch_id=branch[0], employee_id=employee[0] if employee else None)
    if employee:
        as_of = dt.datetime.fromisoformat(payload.get("received_at") or msg["timestamp"])
        for key, index in [("position", 2), ("education", 3)]:
            if key not in detail and employee[index] is not None:
                detail[key] = employee[index]
        for key, index in [("tenure_org_days", 4), ("tenure_position_days", 5)]:
            if key not in detail and employee[index]:
                detail[key] = max(0, (as_of - employee[index]).days)
    if detail.get("tenure_position_days", 0) > detail.get("tenure_org_days", 2147483647):
        raise CsiError("tenure_position_days", "стаж в должности больше стажа в организации")
    msg["meta"] = {**(msg.get("meta") or {}), "csi_meta": payload}
    msg["geo"] = {
        "source": "message",
        "region_code": branch[1],
        "city_id": branch[2],
        "lat": float(branch[3]) if branch[3] is not None else None,
        "lon": float(branch[4]) if branch[4] is not None else None,
    }
    return detail, {**msg["geo"], "vsp_code": payload["branch_code"], "region_okato": branch[5]}


async def persist(conn: psycopg.AsyncConnection[Any], signal_id: int, detail: dict[str, Any]) -> None:
    columns = list(detail)
    await conn.execute(
        "insert into csi_responses (signal_id,"
        + ",".join(columns)
        + ") values (%s,"
        + ",".join("%s" for _ in columns)
        + ")",
        (signal_id, *detail.values()),
    )


async def import_rows(conn: psycopg.AsyncConnection[Any], rows: list[Row]) -> dict[str, Any]:
    from app.ingest.normalizer import IngestError, create_signal, from_csi

    current = await _one(conn, "select current_run_id,sim_time from sim_state where id=1")
    if not current or not current[0]:
        raise IngestError("нет текущего запуска: нужен сид")
    report: dict[str, Any] = {"accepted": 0, "duplicates": 0, "rejected": 0, "items": []}
    async with conn.transaction():
        for row in rows:
            try:
                payload = validate_row(row.value)
                result = await create_signal(
                    conn, from_csi(payload), source_code="csi", run_id=int(current[0]), sim_time=current[1]
                )
                status = result["status"]
                if status == "rejected":
                    raise CsiError("<row>", str(result.get("error") or "ошибка нормализации"))
                report["accepted" if status == "accepted" else "duplicates"] += 1
                report["items"].append({"row": row.number, **result})
            except (CsiError, IngestError) as exc:
                column = exc.column if isinstance(exc, CsiError) else "<row>"
                report["rejected"] += 1
                report["items"].append(
                    {
                        "row": row.number,
                        "status": "rejected",
                        "column": row.columns.get(column, column),
                        "field": column,
                        "error": str(exc),
                    }
                )
                raw = row.value if isinstance(row.value, dict) else {"value": row.value}
                ext = str(raw.get("csi_no") or ("invalid-" + hashlib.sha256(dumps(raw).encode()).hexdigest()[:24]))[
                    :128
                ]
                await conn.execute(
                    "insert into ingest_messages (run_id,source_code,external_id,payload,status,error) values (%s,'csi',%s,%s::jsonb,'rejected',%s) on conflict (source_code,external_id) do update set payload=excluded.payload,error=excluded.error,status='rejected' where ingest_messages.status<>'accepted'",
                    (current[0], ext, dumps(raw), str(exc)[:500]),
                )
    return report
